"""분석 결과 → 엑셀 파일"""
import os
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config as C
from .accounts import ACCOUNTS
from .analyzer import RATIO_SPEC

FMT = {"%": '0.0"%"', "x": '0.00"배"', "days": '0"일"', "amt": '#,##0;[Red]-#,##0', "won": '#,##0"원"'}
NAVY = PatternFill("solid", fgColor="1F3864")
GREY = PatternFill("solid", fgColor="F2F2F2")
LIGHT_FILL = {"🟢": "C6EFCE", "🟡": "FFEB9C", "🔴": "FFC7CE", "🟠": "F8CBAD", "⚪": "EDEDED"}
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
H_FONT = Font(bold=True, color="FFFFFF")


def _amt(x):
    return None if x is None else round(x / C.UNIT)


def _cell(ws, r, c, val, fmt=None, bold=False, fill=None, color=None, wrap=False, border=True):
    cell = ws.cell(row=r, column=c, value="-" if val is None else val)
    if fmt and val is not None:
        cell.number_format = fmt
    if bold or color:
        cell.font = Font(bold=bold, color=color)
    if fill:
        cell.fill = fill
    if border:
        cell.border = THIN
    cell.alignment = Alignment(vertical="center", wrap_text=wrap,
                               horizontal="right" if isinstance(val, (int, float)) else None)
    return cell


def _header(ws, r, labels, start=1):
    for i, h in enumerate(labels):
        c = ws.cell(row=r, column=start + i, value=h)
        c.font, c.fill, c.border = H_FONT, NAVY, THIN
        c.alignment = Alignment(horizontal="center", vertical="center")


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(color="595959")


# ── 시트: 요약 ────────────────────────────────────────
def _summary(wb, rep):
    ws = wb.active
    ws.title = "요약"
    m, years, V, Rt = rep["meta"], rep["years"], rep["V"], rep["R"]
    last = years[-1]
    _title(ws, f"{m['name']}({m['code']}) 재무제표 분석",
           f"기준: {rep['basis']}재무제표 | 기간: FY{years[0]}~FY{last} | 단위: {C.UNIT_LABEL} | 생성일: {m['date']}")
    r = 4
    for light, title, detail in (rep["cs_verdicts"] or []):
        if light != "🟢":
            c = ws.cell(row=r, column=1, value=f"{light} 연결vs별도 경고 — {title}: {detail}")
            c.font = Font(bold=True, color="9C0006")
            c.fill = PatternFill("solid", fgColor=LIGHT_FILL[light])
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            r += 1
    v = rep["valuation"]
    if v["treasury_pct"] and v["treasury_pct"] >= C.TREASURY_WARN_PCT:
        c = ws.cell(row=r, column=1, value=f"⚠️ 자사주 비중 {v['treasury_pct']:.1f}% — 소각 여부 확인 필요 "
                                           "(유통주식 기준 PER이 낮게 보일 수 있음)")
        c.font = Font(bold=True, color="9C5700")
        c.fill = PatternFill("solid", fgColor="FFEB9C")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        r += 1
    r += 1

    ws.cell(row=r, column=1, value="종합 신호").font = Font(bold=True, size=12)
    r += 1
    _header(ws, r, ["구분", "신호", "근거"])
    for cat, light, why in rep["signals"]:
        r += 1
        _cell(ws, r, 1, cat, bold=True)
        _cell(ws, r, 2, light, fill=PatternFill("solid", fgColor=LIGHT_FILL[light]))
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
        _cell(ws, r, 3, why, wrap=True)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    r += 2

    ws.cell(row=r, column=1, value=f"핵심 지표 (FY{last})").font = Font(bold=True, size=12)
    r += 1
    L, VL = Rt[last], V[last]
    kp = [("매출액", _amt(VL.get("revenue")), FMT["amt"]), ("영업이익", _amt(VL.get("op_income")), FMT["amt"]),
          ("지배주주순이익", _amt(VL.get("ni_owner")), FMT["amt"]), ("FCF", _amt(VL.get("fcf")), FMT["amt"]),
          ("영업이익률", L.get("opm"), FMT["%"]), ("ROE", L.get("roe"), FMT["%"]),
          ("ROIC", L.get("roic"), FMT["%"]), ("부채비율", L.get("debt_ratio"), FMT["%"])]
    _header(ws, r, ["지표", "값", "", "지표", "값"])
    for i in range(4):
        r += 1
        for j, (lab, val, f) in enumerate((kp[i], kp[i + 4])):
            _cell(ws, r, 1 + j * 3, lab, bold=True, fill=GREY)
            _cell(ws, r, 2 + j * 3, val, f)
    r += 2

    ws.cell(row=r, column=1, value=f"밸류에이션 (현재가 {v['price']:,}원)").font = Font(bold=True, size=12)
    r += 1
    _header(ws, r, ["지표", "유통주식 기준", "발행주식 기준"])
    fv, iv = v["float"], v["issued"]
    for lab, a, b, f in [("주식수", v["float_shares"], v["issued_shares"], '#,##0"주"'),
                         (f"시가총액({C.UNIT_LABEL})", _amt(fv["mcap"]), _amt(iv["mcap"]), FMT["amt"]),
                         ("EPS", fv["eps"], iv["eps"], FMT["won"]), ("BPS", fv["bps"], iv["bps"], FMT["won"]),
                         ("PER", fv["per"], iv["per"], FMT["x"]), ("PBR", fv["pbr"], iv["pbr"], FMT["x"]),
                         ("EV/EBITDA", fv["ev_ebitda"], iv["ev_ebitda"], FMT["x"])]:
        r += 1
        _cell(ws, r, 1, lab, bold=True, fill=GREY)
        _cell(ws, r, 2, a, f, bold=True)
        _cell(ws, r, 3, b, f)
    for lab, val, f in [("자사주 비중", v["treasury_pct"], FMT["%"]), ("배당성향", v["payout"], FMT["%"]),
                        ("배당수익률(추정)", v["div_yield"], FMT["%"])]:
        r += 1
        _cell(ws, r, 1, lab, bold=True, fill=GREY)
        _cell(ws, r, 2, val, f)
    r += 1
    ws.cell(row=r, column=1, value="※ 배당은 해당 연도 현금흐름표의 배당금 지급액(대부분 전년도 결산 배당) 기준 추정치").font = \
        Font(italic=True, color="7F7F7F", size=9)
    r += 2

    ws.cell(row=r, column=1, value="해석 코멘트").font = Font(bold=True, size=12)
    for line in rep["comments"]:
        r += 1
        c = ws.cell(row=r, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 32
    if rep["notes"]:
        r += 2
        ws.cell(row=r, column=1, value="데이터 메모").font = Font(bold=True, size=12)
        for nt, ys in rep["notes"].items():
            r += 1
            ws.cell(row=r, column=1, value=f"· {nt} (FY{', FY'.join(map(str, ys))})").font = Font(color="595959")
    _widths(ws, [20, 16, 16, 20, 16, 30])


# ── 시트: 원본데이터 ──────────────────────────────────
DERIVED = [("debt", "차입금 합계(이자부부채)"), ("cash_like", "현금성자산+단기금융상품"),
           ("net_debt", "순차입금"), ("capex", "CAPEX(유형+무형 취득)"), ("da", "감가상각비 합계"),
           ("ebitda", "EBITDA"), ("fcf", "FCF")]


def _raw(wb, title, years, V, names):
    ws = wb.create_sheet(title)
    _title(ws, title, f"단위: {C.UNIT_LABEL} | 출처: DART 사업보고서")
    _header(ws, 4, ["항목"] + [f"FY{y}" for y in years] + ["사용한 DART 계정명(최근연도)"])
    r = 4
    for key, (label, *_rest) in ACCOUNTS.items():
        r += 1
        _cell(ws, r, 1, label, bold=True, fill=GREY)
        for j, y in enumerate(years):
            _cell(ws, r, 2 + j, _amt(V[y].get(key)), FMT["amt"])
        _cell(ws, r, 2 + len(years), names.get(key) or "(매칭 실패)",
              color=None if names.get(key) else "C00000")
    r += 1
    _cell(ws, r, 1, "— 계산 항목 —", bold=True, border=False)
    for key, label in DERIVED:
        r += 1
        _cell(ws, r, 1, label, bold=True, fill=GREY)
        for j, y in enumerate(years):
            _cell(ws, r, 2 + j, _amt(V[y].get(key)), FMT["amt"])
        if key == "debt":
            _cell(ws, r, 2 + len(years), names.get("debt", ""), wrap=True)
    _widths(ws, [26] + [13] * len(years) + [40])
    ws.freeze_panes = "B5"


# ── 시트: 재무비율 ────────────────────────────────────
def _ratios(wb, rep):
    ws = wb.create_sheet("재무비율")
    years, Rt = rep["years"], rep["R"]
    _title(ws, "재무비율", f"{rep['basis']} 기준 | 금액 단위: {C.UNIT_LABEL}")
    _header(ws, 4, ["구분", "지표"] + [f"FY{y}" for y in years] + [f"{len(years)}년 평균/CAGR"])
    r, prev = 4, None
    cagr_map = {"rev_g": "revenue", "op_g": "op_income", "ni_g": "ni_owner"}
    for cat, key, label, kind in RATIO_SPEC:
        r += 1
        _cell(ws, r, 1, cat if cat != prev else "", bold=True, fill=GREY)
        prev = cat
        _cell(ws, r, 2, label)
        vals = []
        for j, y in enumerate(years):
            x = Rt[y].get(key)
            vals.append(x)
            _cell(ws, r, 3 + j, _amt(x) if kind == "amt" else x, FMT[kind])
        if key in cagr_map:
            agg = rep["cagr"].get(cagr_map[key])
            _cell(ws, r, 3 + len(years), agg, '"CAGR "0.0"%"')
        else:
            xs = [x for x in vals if x is not None]
            agg = sum(xs) / len(xs) if xs else None
            _cell(ws, r, 3 + len(years), _amt(agg) if kind == "amt" else agg, FMT[kind])
    _widths(ws, [10, 26] + [12] * len(years) + [16])
    ws.freeze_panes = "C5"


# ── 시트: 연결vs별도 ──────────────────────────────────
def _cons_sep(wb, rep):
    ws = wb.create_sheet("연결vs별도")
    rows = rep["cs_rows"]
    ys = sorted(rows)
    _title(ws, "연결 vs 별도 비교", f"금액 단위: {C.UNIT_LABEL}")
    r = 4
    for light, title, detail in rep["cs_verdicts"]:
        _cell(ws, r, 1, f"{light} {title}", bold=True, fill=PatternFill("solid", fgColor=LIGHT_FILL[light]))
        _cell(ws, r, 2, detail, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(ys) + 1)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    _header(ws, r, ["항목"] + [f"FY{y}" for y in ys])
    spec = [("매출액 · 연결", "c_rev", "amt"), ("매출액 · 별도", "s_rev", "amt"),
            ("영업이익 · 연결", "c_op", "amt"), ("영업이익 · 별도", "s_op", "amt"),
            ("순이익 · 연결(지배)", "c_ni", "amt"), ("순이익 · 별도", "s_ni", "amt"),
            ("영업CF · 연결", "c_cfo", "amt"), ("영업CF · 별도", "s_cfo", "amt"),
            ("① 자회사 이익 의존도", "dep", "%"), ("② 별도 적자·연결 흑자", "sep_loss", "bool"),
            ("③ 별도 순이익 중 배당·지분법 비중", "div_dep", "%"), ("④ 별도 영업CF ÷ 연결 영업CF", "cfo_ratio", "x"),
            ("⑤ 비지배지분 순이익 비중", "nci", "%"), ("⑥ 부채비율 차이(연결−별도, %p)", "debt_gap", "pp"),
            ("⑦ 연결 영업이익이 별도보다 30%↓", "cons_worse", "bool")]
    for label, key, kind in spec:
        r += 1
        _cell(ws, r, 1, label, bold=True, fill=GREY)
        for j, y in enumerate(ys):
            x = rows[y].get(key)
            if kind == "bool":
                cc = _cell(ws, r, 2 + j, "해당" if x else "-",
                           fill=PatternFill("solid", fgColor="FFC7CE") if x else None)
                cc.alignment = Alignment(horizontal="center")
            elif kind == "amt":
                _cell(ws, r, 2 + j, _amt(x), FMT["amt"])
            elif kind == "pp":
                _cell(ws, r, 2 + j, x, '+0.0"%p";-0.0"%p"')
            else:
                _cell(ws, r, 2 + j, x, FMT[kind])
    r += 2
    ws.cell(row=r, column=1, value="※ ①·③은 별도 이익이 적자이거나 아주 작으면 100%를 넘을 수 있습니다(자회사가 모회사 손실까지 메운 상태).").font = \
        Font(italic=True, color="7F7F7F", size=9)
    r += 1
    ws.cell(row=r, column=1, value="※ ③은 배당금수익·지분법이익이 손익계산서 본문에 있을 때만 계산됩니다(주석에만 있으면 '-').").font = \
        Font(italic=True, color="7F7F7F", size=9)
    _widths(ws, [34] + [14] * len(ys))


# ── 시트: 경고신호 ────────────────────────────────────
def _warnings(wb, rep):
    ws = wb.create_sheet("경고신호")
    _title(ws, "경고 신호", f"{rep['basis']} 기준 자동 체크")
    _header(ws, 4, ["연도", "구분", "근거"])
    r = 4
    if not rep["warnings"]:
        _cell(ws, 5, 1, "-")
        _cell(ws, 5, 2, "해당 없음")
    for w in rep["warnings"]:
        r += 1
        _cell(ws, r, 1, f"FY{w['year']}")
        _cell(ws, r, 2, f"🚩 {w['kind']}", bold=True)
        _cell(ws, r, 3, w["detail"], wrap=True)
    _widths(ws, [10, 24, 70])


# ── 시트: 차트 ────────────────────────────────────────
def _charts(wb, rep):
    ws = wb.create_sheet("차트")
    years, V, Rt = rep["years"], rep["V"], rep["R"]
    _title(ws, "추이 차트", f"금액 단위: {C.UNIT_LABEL}")
    cols = [("연도", lambda y: f"FY{y}"),
            ("매출액", lambda y: _amt(V[y].get("revenue"))), ("영업이익", lambda y: _amt(V[y].get("op_income"))),
            ("지배주주순이익", lambda y: _amt(V[y].get("ni_owner"))),
            ("매출총이익률", lambda y: Rt[y].get("gpm")), ("영업이익률", lambda y: Rt[y].get("opm")),
            ("순이익률", lambda y: Rt[y].get("npm")),
            ("영업CF", lambda y: _amt(V[y].get("cfo"))), ("CAPEX", lambda y: _amt(V[y].get("capex"))),
            ("FCF", lambda y: _amt(V[y].get("fcf"))),
            ("ROE", lambda y: Rt[y].get("roe")), ("ROIC", lambda y: Rt[y].get("roic"))]
    _header(ws, 4, [c[0] for c in cols])
    for i, y in enumerate(years):
        for j, (_, fn) in enumerate(cols):
            val = fn(y)
            c = ws.cell(row=5 + i, column=1 + j, value=val)
            c.border = THIN
            if j > 0:
                c.number_format = '0.0' if j in (4, 5, 6, 10, 11) else '#,##0;[Red]-#,##0'
    last_row = 4 + len(years)
    cats = Reference(ws, min_col=1, min_row=5, max_row=last_row)

    def chart(cls, title, c1, c2, anchor, ytitle):
        ch = cls()
        ch.title, ch.height, ch.width = title, 7.5, 15
        ch.y_axis.title = ytitle
        ch.add_data(Reference(ws, min_col=c1, max_col=c2, min_row=4, max_row=last_row), titles_from_data=True)
        ch.set_categories(cats)
        ws.add_chart(ch, anchor)

    chart(BarChart, f"매출·이익 추이({C.UNIT_LABEL})", 2, 4, f"A{last_row + 3}", C.UNIT_LABEL)
    chart(LineChart, "마진 추이(%)", 5, 7, f"H{last_row + 3}", "%")
    chart(BarChart, f"현금흐름({C.UNIT_LABEL})", 8, 10, f"A{last_row + 19}", C.UNIT_LABEL)
    chart(LineChart, "ROE·ROIC(%)", 11, 12, f"H{last_row + 19}", "%")
    _widths(ws, [10] + [13] * 11)


def write_excel(rep, out_dir="."):
    wb = Workbook()
    _summary(wb, rep)
    _ratios(wb, rep)
    if rep["sep"]:
        _cons_sep(wb, rep)
    _warnings(wb, rep)
    _charts(wb, rep)
    _raw(wb, f"원본데이터({rep['basis']})", rep["years"], rep["V"], rep["names"])
    if rep["sep"]:
        _raw(wb, "원본데이터(별도)", rep["sep"]["years"], rep["sep"]["V"], rep["sep_names"])
    m = rep["meta"]
    safe = re.sub(r'[\\/:*?"<>|]', "", m["name"])
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"{safe}_{m['code']}_재무분석_{m['date'].replace('-', '')}.xlsx")
    wb.save(path)
    return path
